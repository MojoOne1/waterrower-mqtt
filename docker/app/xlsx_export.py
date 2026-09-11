"""Excel workbooks: one per session, and one across all sessions."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=13)
_DATETIME = "yyyy-mm-dd hh:mm"
_DURATION = "[h]:mm:ss"
_SPLIT = "m:ss"


def _clock(ts):
    return datetime.fromtimestamp(ts) if ts else None


def _days(seconds):
    """Excel keeps durations as a fraction of a day, so they stay computable."""
    return (seconds or 0) / 86400


def _split(speed_ms):
    return 500 / speed_ms / 86400 if speed_ms else None


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_session(session: dict, samples: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    speeds = [s["speed_ms"] for s in samples if s.get("speed_ms") is not None]
    distance = session.get("distance_m") or 0
    strokes = session.get("strokes") or 0

    ws["A1"] = "WaterRower session"
    ws["A1"].font = _TITLE
    rows = [
        ("Session ID", session["session_id"], None),
        ("Start", _clock(session.get("started_at")), _DATETIME),
        ("End", _clock(session.get("ended_at")), _DATETIME),
        ("Distance", distance, '0 "m"'),
        ("Duration", _days(session.get("duration_s")), _DURATION),
        ("Strokes", strokes, "0"),
        ("Avg speed", session.get("avg_speed_ms") or 0, '0.00 "m/s"'),
        ("Avg 500 m split", _split(session.get("avg_speed_ms")), _SPLIT),
        ("Peak speed", max(speeds) if speeds else 0, '0.00 "m/s"'),
        ("Avg stroke rate", session.get("avg_spm") or 0, '0.0 "spm"'),
        ("Metres per stroke", distance / strokes if strokes else 0, '0.0 "m"'),
        ("Samples", len(samples), "0"),
    ]
    for i, (label, value, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = _BOLD
        cell = ws.cell(row=i, column=2, value=value)
        if fmt:
            cell.number_format = fmt
    _widths(ws, [20, 22])

    sh = wb.create_sheet("Samples")
    sh.append(["Time (s)", "Clock", "Distance (m)", "Speed (m/s)", "500 m split",
               "Stroke rate (spm)", "Strokes", "Watts"])
    for cell in sh[1]:
        cell.font = _BOLD
    sh.freeze_panes = "A2"
    t0 = samples[0]["ts"] if samples else 0
    for s in samples:
        sh.append([
            round(s["ts"] - t0, 1), _clock(s["ts"]), s.get("distance_m"),
            s.get("speed_ms"), _split(s.get("speed_ms")), s.get("stroke_rate"),
            s.get("strokes"), s.get("watts"),
        ])
    for row in sh.iter_rows(min_row=2, min_col=2, max_col=5):
        row[0].number_format = "hh:mm:ss"
        row[3].number_format = _SPLIT
    _widths(sh, [10, 12, 14, 12, 12, 18, 10, 8])

    # One chart per measure: two y-scales in a single plot make unrelated
    # curves look correlated, so speed and stroke rate get their own.
    if len(samples) > 1:
        last = len(samples) + 1
        cats = Reference(sh, min_col=1, min_row=2, max_row=last)
        for col, title, unit, anchor in ((4, "Speed", "m/s", "D3"), (6, "Stroke rate", "spm", "D21")):
            chart = LineChart()
            chart.title = title
            chart.height, chart.width = 8, 26
            chart.y_axis.title = unit
            chart.x_axis.title = "Time (s)"
            chart.add_data(Reference(sh, min_col=col, min_row=1, max_row=last), titles_from_data=True)
            chart.set_categories(cats)
            chart.legend = None      # single series, the title names it
            for series in chart.series:
                series.smooth = False
            ws.add_chart(chart, anchor)

    return _save(wb)


def build_overview(sessions: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.append(["Date", "Session ID", "Distance (m)", "Duration", "500 m split",
               "Avg speed (m/s)", "Avg stroke rate (spm)", "Strokes", "Metres per stroke"])
    for cell in ws[1]:
        cell.font = _BOLD
    ws.freeze_panes = "A2"

    chrono = sorted(sessions, key=lambda s: s.get("started_at") or 0)
    for s in chrono:
        distance = s.get("distance_m") or 0
        strokes = s.get("strokes") or 0
        ws.append([
            _clock(s.get("started_at")), s["session_id"], distance,
            _days(s.get("duration_s")), _split(s.get("avg_speed_ms")),
            s.get("avg_speed_ms"), s.get("avg_spm"), strokes,
            round(distance / strokes, 1) if strokes else None,
        ])
    last = len(chrono) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        row[0].number_format = _DATETIME
        row[3].number_format = _DURATION
        row[4].number_format = _SPLIT
        row[5].number_format = "0.00"
        row[6].number_format = "0.0"

    # Live formulas, so the totals still add up after rows are edited.
    total = last + 2
    ws.cell(row=total, column=2, value="Total").font = _BOLD
    for col, fmt in ((3, "0"), (4, _DURATION), (8, "0")):
        letter = get_column_letter(col)
        cell = ws.cell(row=total, column=col, value=f"=SUM({letter}2:{letter}{last})")
        cell.number_format = fmt
        cell.font = _BOLD
    _widths(ws, [18, 18, 14, 12, 12, 16, 20, 10, 18])

    if len(chrono) > 1:
        chart = BarChart()
        chart.title = "Distance per session"
        chart.y_axis.title = "m"
        chart.height, chart.width = 9, 26
        chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        chart.legend = None
        ws.add_chart(chart, f"A{total + 2}")

    return _save(wb)
