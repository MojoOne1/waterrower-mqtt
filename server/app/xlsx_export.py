"""Excel workbooks: one per session, and one across the whole arena.

The tracker has the same idea for a single machine. Here every row carries
an athlete, which is the point - the overview is what you open when you want
to see the three of you side by side in something you can sort and filter.

Samples are keyed by `t`, the device's own elapsed seconds, rather than by a
wall clock: it is what the charts are drawn against and what survives a
backfill, where no arrival time was ever recorded.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=13)
_DATETIME = "yyyy-mm-dd hh:mm"
_DURATION = "[h]:mm:ss"
_SPLIT = "m:ss"


def _clock(ts):
    return datetime.fromtimestamp(ts) if ts else None


def _days(seconds):
    """Excel keeps durations as a fraction of a day, so they stay computable."""
    return (seconds or 0) / 86400


def _split(speed_ms):
    return 500 / speed_ms / 86400 if speed_ms else None


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_session(session: dict, samples: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    speeds = [s["speed_ms"] for s in samples if s.get("speed_ms") is not None]
    watts = [s["watts"] for s in samples if s.get("watts") is not None]
    distance = session.get("distance_m") or 0
    strokes = session.get("strokes") or 0

    ws["A1"] = "WaterRower session"
    ws["A1"].font = _TITLE
    rows = [
        ("Athlete", session.get("display_name"), None),
        ("Session", session.get("remote_id"), None),
        ("Start", _clock(session.get("started_at")), _DATETIME),
        ("End", _clock(session.get("ended_at")), _DATETIME),
        ("Distance", distance, '0 "m"'),
        ("Duration", _days(session.get("duration_s")), _DURATION),
        ("Strokes", strokes, "0"),
        ("Avg speed", session.get("avg_speed_ms") or 0, '0.00 "m/s"'),
        ("Avg 500 m split", _split(session.get("avg_speed_ms")), _SPLIT),
        ("Peak speed", max(speeds) if speeds else 0, '0.00 "m/s"'),
        ("Avg stroke rate", session.get("avg_spm") or 0, '0.0 "spm"'),
        ("Avg watts", session.get("avg_watts") or 0, '0 "W"'),
        ("Peak watts", max(watts) if watts else 0, '0 "W"'),
        ("Metres per stroke", distance / strokes if strokes else 0, '0.0 "m"'),
        ("Part of a race", "yes" if session.get("race_id") else "no", None),
        ("Samples", len(samples), "0"),
    ]
    for i, (label, value, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = _BOLD
        cell = ws.cell(row=i, column=2, value=value)
        if fmt:
            cell.number_format = fmt
    _widths(ws, [20, 24])

    sh = wb.create_sheet("Samples")
    sh.append(["Time (s)", "Distance (m)", "Speed (m/s)", "500 m split",
               "Stroke rate (spm)", "Strokes", "Watts"])
    for cell in sh[1]:
        cell.font = _BOLD
    sh.freeze_panes = "A2"
    for s in samples:
        sh.append([
            s.get("t"), s.get("distance_m"), s.get("speed_ms"),
            _split(s.get("speed_ms")), s.get("stroke_rate"),
            s.get("strokes"), s.get("watts"),
        ])
    for row in sh.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = _SPLIT
    _widths(sh, [10, 14, 12, 12, 18, 10, 8])

    # One chart per measure: two y-scales in a single plot make unrelated
    # curves look correlated, so speed and stroke rate get their own.
    if len(samples) > 1:
        last = len(samples) + 1
        cats = Reference(sh, min_col=1, min_row=2, max_row=last)
        for col, title, unit, anchor in ((3, "Speed", "m/s", "D3"),
                                         (5, "Stroke rate", "spm", "D21")):
            chart = LineChart()
            chart.title = title
            chart.height, chart.width = 8, 26
            chart.y_axis.title = unit
            chart.x_axis.title = "Time (s)"
            chart.add_data(Reference(sh, min_col=col, min_row=1, max_row=last),
                           titles_from_data=True)
            chart.set_categories(cats)
            chart.legend = None      # single series, the title names it
            for series in chart.series:
                series.smooth = False
            ws.add_chart(chart, anchor)

    return _save(wb)


def build_overview(sessions: list[dict], totals: list[dict]) -> bytes:
    """Every session as a row, plus a sheet of per-athlete totals."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.append(["Date", "Athlete", "Session", "Distance (m)", "Duration",
               "500 m split", "Avg speed (m/s)", "Avg stroke rate (spm)",
               "Avg watts", "Strokes", "Metres per stroke", "Race"])
    for cell in ws[1]:
        cell.font = _BOLD
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:L1"     # sort and filter by athlete, which is the point

    chrono = sorted(sessions, key=lambda s: s.get("started_at") or 0)
    for s in chrono:
        distance = s.get("distance_m") or 0
        strokes = s.get("strokes") or 0
        ws.append([
            _clock(s.get("started_at")), s.get("display_name"), s.get("remote_id"),
            distance, _days(s.get("duration_s")), _split(s.get("avg_speed_ms")),
            s.get("avg_speed_ms"), s.get("avg_spm"), s.get("avg_watts"), strokes,
            round(distance / strokes, 1) if strokes else None,
            "yes" if s.get("race_id") else "",
        ])
    last = len(chrono) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        row[0].number_format = _DATETIME
        row[4].number_format = _DURATION
        row[5].number_format = _SPLIT
        row[6].number_format = "0.00"
        row[7].number_format = "0.0"
        row[8].number_format = "0"
    _widths(ws, [18, 16, 18, 14, 12, 12, 16, 20, 12, 10, 18, 8])

    tw = wb.create_sheet("Totals")
    tw.append(["Athlete", "Sessions", "Distance (m)", "Duration", "Strokes"])
    for cell in tw[1]:
        cell.font = _BOLD
    for row in totals:
        tw.append([row.get("display_name"), row.get("sessions"),
                   row.get("distance_m"), _days(row.get("duration_s")),
                   row.get("strokes")])
    for r in tw.iter_rows(min_row=2, min_col=4, max_col=4):
        r[0].number_format = _DURATION
    _widths(tw, [16, 12, 14, 12, 10])

    if len(totals) > 1:
        chart = BarChart()
        chart.title = "Distance per athlete"
        chart.y_axis.title = "m"
        chart.height, chart.width = 9, 20
        rows = len(totals) + 1
        chart.add_data(Reference(tw, min_col=3, min_row=1, max_row=rows),
                       titles_from_data=True)
        chart.set_categories(Reference(tw, min_col=1, min_row=2, max_row=rows))
        chart.legend = None
        tw.add_chart(chart, "G2")

    return _save(wb)
